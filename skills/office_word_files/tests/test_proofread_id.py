import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
TOOL = ROOT / "proofread_id.py"
VENV_PY = ROOT / ".venv-mcp" / "bin" / "python"
PY = str(VENV_PY) if VENV_PY.exists() else sys.executable


def _sample(path: Path) -> None:
    doc = Document()
    doc.add_heading("JUDUL", level=1)
    doc.add_paragraph("Ini kalimat  dengan spasi ganda dan kata kata ganda.")
    doc.add_paragraph("Pembayaran di proses manual, lihat website kami.")
    doc.add_paragraph("Skor 4,46 valid. [PERLU REFERENSI: contoh]")
    doc.save(path)


def test_proofread_menemukan_kesalahan(tmp_path: Path):
    f = tmp_path / "uji.docx"
    _sample(f)
    r = subprocess.run([PY, str(TOOL), str(f)], capture_output=True, text=True)
    assert r.returncode == 1, r.stdout
    for kunci in ("spasi ganda", "kata ganda", "dirangkai", "website",
                  "penanda draft"):
        assert kunci in r.stdout, f"hilang: {kunci}\n{r.stdout}"


def test_proofread_bersih(tmp_path: Path):
    f = tmp_path / "bersih.docx"
    doc = Document()
    doc.add_heading("JUDUL", level=1)
    doc.add_paragraph("Pembayaran diproses manual dan diterbitkan otomatis.")
    doc.save(f)
    r = subprocess.run([PY, str(TOOL), str(f)], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout
    assert "BERSIH" in r.stdout


def test_proofread_cek_jurnal(tmp_path: Path):
    f = tmp_path / "jurnal.docx"
    doc = Document()
    doc.add_heading("ABSTRAK", level=1)
    doc.add_paragraph("kata " * 100)  # 100 kata -> di bawah 150
    doc.add_paragraph("Kata kunci: satu, dua")
    doc.save(f)
    r = subprocess.run([PY, str(TOOL), str(f)], capture_output=True, text=True)
    assert r.returncode == 1, r.stdout
    assert "150–250" in r.stdout and "3–6" in r.stdout
