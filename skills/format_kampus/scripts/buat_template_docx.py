#!/usr/bin/env python3
"""Hasilkan template acuan DOCX dari profil format kampus (YAML).

Pemakaian: python3 buat_template_docx.py profil-kampus.yaml template-kampus.docx

Template yang dihasilkan dapat dipakai sebagai --reference-doc pandoc atau basis
dokumen python-docx. Semua slot font (ascii, hAnsi, eastAsia, cs) diisi literal
dan atribut font tema dibuang, agar Microsoft Word tidak mengganti font dengan
font tema (mis. Aptos) — penyebab klasik "di Word bukan Times New Roman".
"""
import re
import shutil
import sys
import zipfile

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

RATA = {'tengah': WD_ALIGN_PARAGRAPH.CENTER, 'kiri': WD_ALIGN_PARAGRAPH.LEFT,
        'kanan': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
        'rata_kiri_kanan': WD_ALIGN_PARAGRAPH.JUSTIFY}
KERTAS = {'A4': (21.0, 29.7), 'Letter': (21.59, 27.94), 'F4': (21.59, 33.02)}


def kunci_font(style, nama):
    """Isi semua slot rFonts secara literal dan hapus atribut tema."""
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.insert(0, rf)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(attr)) is not None:
            del rf.attrib[qn(attr)]
    for slot in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(slot), nama)


def terapkan_style(style, nama_font, pt, tebal=None, kapital=False, rata=None, spasi=None):
    f = style.font
    f.name = nama_font
    f.size = Pt(pt)
    if tebal is not None:
        f.bold = tebal
    f.all_caps = kapital
    kunci_font(style, nama_font)
    pf = style.paragraph_format
    if rata:
        pf.alignment = RATA.get(rata)
    if spasi:
        pf.line_spacing = float(spasi)



def _usage(msg: str = "", kode: int = 0) -> None:
    if msg:
        print(msg, file=sys.stderr)
    print(__doc__.strip(), file=sys.stderr if kode else sys.stdout)
    raise SystemExit(kode)


def _butuh(n: int, contoh: str) -> None:
    if len(sys.argv) - 1 < n:
        _usage(f"Kurang argumen. Contoh: {contoh}", 2)

def main():
    if '-h' in sys.argv or '--help' in sys.argv:
        _usage()
    _butuh(2, 'python3 buat_template_docx.py profil-kampus.yaml template-kampus.docx')
    profil = yaml.safe_load(open(sys.argv[1], encoding='utf8'))
    keluar = sys.argv[2]
    dok = profil.get('dokumen') or {}
    mar = dok.get('margin_cm') or {}
    fon = profil.get('font') or {}
    spa = profil.get('spasi') or {}
    nama_font = fon.get('nama') or 'Times New Roman'

    d = Document()
    s = d.sections[0]
    w, h = KERTAS.get(dok.get('ukuran_kertas') or 'A4', KERTAS['A4'])
    s.page_width, s.page_height = Cm(w), Cm(h)
    s.left_margin = Cm(mar.get('kiri', 4))
    s.top_margin = Cm(mar.get('atas', 3))
    s.right_margin = Cm(mar.get('kanan', 3))
    s.bottom_margin = Cm(mar.get('bawah', 3))

    jb = fon.get('judul_bab') or {}
    sb = fon.get('subbab') or {}
    ss = fon.get('sub_subbab') or {}
    terapkan_style(d.styles['Normal'], nama_font, fon.get('isi_pt', 12),
                   spasi=spa.get('isi', 1.5))
    d.styles['Normal'].paragraph_format.widow_control = True
    terapkan_style(d.styles['Heading 1'], nama_font, jb.get('pt', 14),
                   tebal=jb.get('tebal', True), kapital=jb.get('kapital', True),
                   rata=jb.get('rata', 'tengah'), spasi=spa.get('judul', 1.5))
    terapkan_style(d.styles['Heading 2'], nama_font, sb.get('pt', 12),
                   tebal=sb.get('tebal', True), spasi=spa.get('judul', 1.5))
    terapkan_style(d.styles['Heading 3'], nama_font, ss.get('pt', 12),
                   tebal=ss.get('tebal', True), spasi=spa.get('judul', 1.5))
    # kerapian: judul tidak yatim di kaki halaman
    for nama_style in ('Heading 1', 'Heading 2', 'Heading 3'):
        pf = d.styles[nama_style].paragraph_format
        pf.keep_with_next = True
        pf.keep_together = True
        pf.widow_control = True
    # style daftar isi: titik pemimpin ke tab stop kanan selebar teks
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    lebar_teks_twips = int((w - float(mar.get('kiri', 4)) - float(mar.get('kanan', 3))) * 567)
    for tingkat, indent_cm, tebal in (('TOC 1', 0.0, True), ('TOC 2', 0.6, False), ('TOC 3', 1.2, False)):
        try:
            st = d.styles[tingkat]
        except KeyError:
            st = d.styles.add_style(tingkat, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = d.styles['Normal']
        terapkan_style(st, nama_font, fon.get('isi_pt', 12), tebal=tebal, spasi=1.0)
        pf = st.paragraph_format
        pf.left_indent = Cm(indent_cm)
        pf.space_after = Pt(2)
        ppr = st.element.get_or_add_pPr()
        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'right'); tab.set(qn('w:leader'), 'dot'); tab.set(qn('w:pos'), str(lebar_teks_twips))
        tabs.append(tab); ppr.append(tabs)

    d.add_heading('Judul Bab Contoh', level=1)
    d.add_heading('Subbab Contoh', level=2)
    d.add_paragraph('Paragraf contoh sesuai profil kampus. Hapus isi contoh ini '
                    'saat template dipakai sebagai acuan.')
    d.save(keluar)
    bersihkan_paket(keluar, nama_font)
    print('template tersimpan:', keluar)


def bersihkan_paket(path, nama_font):
    """Terapkan kunci font ke SELURUH paket: template bawaan membawa atribut tema
    pada puluhan style laten; Word memprioritaskan tema, jadi semua harus dibuang
    dan semua slot rFonts diisi literal. Typeface latin tema juga diganti."""
    def fix_rfonts(xml):
        def repl(m):
            tag = m.group(0)
            self_close = tag.rstrip().endswith('/>')
            inner = tag[len('<w:rFonts'):].rstrip()
            if self_close:
                inner = inner[:-1].rstrip()
            attrs = dict(re.findall(r'([\w:]+)="([^"]*)"', inner))
            for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
                attrs.pop(a, None)
            for slot in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                attrs[slot] = nama_font
            urut = ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs')
            sisa = [(k, v) for k, v in attrs.items() if k not in urut]
            bagian = [f'{k}="{attrs[k]}"' for k in urut] + [f'{k}="{v}"' for k, v in sisa]
            return '<w:rFonts ' + ' '.join(bagian) + ('/>' if self_close else '>')
        return re.sub(r'<w:rFonts[^>]*?/?>', repl, xml)

    zin = zipfile.ZipFile(path)
    tmp = path + '.tmp'
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.startswith('word/') and item.filename.endswith('.xml'):
            xml = fix_rfonts(data.decode('utf8'))
            if item.filename.endswith('theme1.xml'):
                xml = re.sub(r'<a:latin typeface="[^"]*"', f'<a:latin typeface="{nama_font}"', xml)
            data = xml.encode('utf8')
        zout.writestr(item, data)
    zout.close(); zin.close()
    shutil.move(tmp, path)


if __name__ == '__main__':
    main()
